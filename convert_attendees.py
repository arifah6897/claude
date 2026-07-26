#!/usr/bin/env python3
"""
Convert an Eventbrite-style "wide" attendee export (one column per
checkbox option, marked with 'x') into the flat CSV shape brunch_matcher.py
expects (one comma-joined cell per multi-select question).

Usage:
    python convert_attendees.py attendees.xlsx attendees.csv
"""

import csv
import html
import sys

import openpyxl


def clean(s):
    return html.unescape(s) if s else s


def clean_note(s):
    s = clean(s).strip() if s else ""
    return "" if s.upper() in ("NIL", "NA", "N/A", "-", "NONE") else s


def convert(src_path, out_path):
    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    header2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]

    # A labeled header cell in row 1 starts a group; None cells that follow
    # belong to it until the next labeled cell. The option text for each
    # column in a group lives in row 2.
    groups = {}
    current_label = None
    current_cols = []
    for idx, label in enumerate(header1):
        if label is not None:
            if current_label is not None:
                groups[current_label] = current_cols
            current_label = label
            current_cols = [idx]
        elif current_label is not None:
            current_cols.append(idx)
    if current_label is not None:
        groups[current_label] = current_cols

    goals_cols = groups["What are you hoping to find through Brunch Buddies?"]
    topics_cols = groups["What topics energise you?"]
    name_first = header1.index("First Name")
    name_last = header1.index("Last Name")
    email_col = header1.index("Email") if "Email" in header1 else None
    stage_col = header1.index("Career Stage")
    anything_col = header1.index("Anything else you'd want us to know?")
    notes_col = header1.index("Notes")

    rows = []
    for r in range(3, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c + 1).value for c in range(ws.max_column)]
        if not any(vals):
            continue
        name = f"{vals[name_first] or ''} {vals[name_last] or ''}".strip()
        if not name:
            continue
        goals = [clean(header2[c]) for c in goals_cols if vals[c] == "x"]
        topics = [clean(header2[c]) for c in topics_cols if vals[c] == "x"]
        anything = clean_note(vals[anything_col])
        orgnote = clean_note(vals[notes_col])
        notes = anything
        if orgnote:
            notes = f"{notes} [organizer note: {orgnote}]" if notes else f"[organizer note: {orgnote}]"
        rows.append([
            clean(name),
            (vals[email_col] or "") if email_col is not None else "",
            clean(vals[stage_col]) or "",
            ", ".join(goals),
            ", ".join(topics),
            notes,
        ])

    with open(out_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([
            "Name", "Email", "What stage are you at in your career?",
            "What are you hoping to find through Brunch Buddies?",
            "What topics energise you?",
            "Anything else you'd want us to know?",
        ])
        w.writerows(rows)

    print(f"Wrote {len(rows)} attendees to {out_path}")
    return len(rows)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python convert_attendees.py <input.xlsx> <output.csv>", file=sys.stderr)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
